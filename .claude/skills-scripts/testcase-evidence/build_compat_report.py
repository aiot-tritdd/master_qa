#!/usr/bin/env python3
"""build_compat_report.py — compat.results.json -> <Case>.compat.xlsx (Ma trận + Findings).
Oracle = WCAG 1.4.10 Reflow + parity giữa engine + không lỗi JS. CƠ KHÍ, không reasoning.

Sheet "Ma trận" là thứ đáng xem nhất: engine × viewport. Ô 未実施 phải HIỆN RÕ lý do —
tuyệt đối không để người đọc tưởng 未実施 là "chắc ổn".
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage

SEV_ORDER = {"High": 0, "Medium": 1, "Low": 2}
SEV_FILL = {"High": "C0392B", "Medium": "E67E22", "Low": "F1C40F"}
V_FILL = {"PASS": "27AE60", "FAIL": "C0392B", "未実施": "7F8C8D"}
HEADERS = ["Màn", "Họ", "Engine", "Viewport", "Mức", "Nơi", "Chi tiết (quan sát được)", "Cách fix", "File ảnh", "Evidence"]


def build(results_path, out_path):
    with open(results_path, encoding="utf-8") as f:
        data = json.load(f)
    base = os.path.dirname(os.path.abspath(results_path))
    top = Alignment(wrap_text=True, vertical="top")
    ctr = Alignment(horizontal="center", vertical="center")
    meta = data.get("meta", {})

    # ── Sheet 1: MA TRẬN engine × viewport ──
    wb = Workbook(); ws = wb.active; ws.title = "Ma trận"
    ws.cell(1, 1, f"COMPATIBILITY · {meta.get('case','')} · {meta.get('date','')} · {meta.get('tester','')}").font = Font(bold=True)
    ws.cell(2, 1, "Oracle: WCAG 1.4.10 Reflow (không cuộn ngang @320px) · parity affordance giữa engine · 0 lỗi JS uncaught").font = Font(italic=True, size=9)

    cells = data.get("cells", [])
    engines = sorted({c["browser"] for c in cells})
    vps = [v for v in ["mobile-320", "mobile-390", "tablet-768", "desktop-1280"] if any(c["viewport"] == v for c in cells)]
    vps += sorted({c["viewport"] for c in cells} - set(vps))

    r0 = 4
    ws.cell(r0, 1, "Engine \\ Viewport").font = Font(bold=True, color="FFFFFF")
    ws.cell(r0, 1).fill = PatternFill("solid", fgColor="2C3E50")
    for j, v in enumerate(vps, 2):
        c = ws.cell(r0, j, v); c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2C3E50"); c.alignment = ctr
    for i, e in enumerate(engines, 1):
        rr = r0 + i
        ec = ws.cell(rr, 1, e); ec.font = Font(bold=True)
        for j, v in enumerate(vps, 2):
            cell_data = next((c for c in cells if c["browser"] == e and c["viewport"] == v), None)
            val = cell_data.get("result", "") if cell_data else ""
            cc = ws.cell(rr, j, val)
            cc.alignment = ctr
            if val in V_FILL:
                cc.fill = PatternFill("solid", fgColor=V_FILL[val]); cc.font = Font(bold=True, color="FFFFFF")
            note = (cell_data or {}).get("note", "")
            if note:
                cc.comment = None
                ws.cell(rr, len(vps) + 2, note).alignment = top   # lý do hiện NGAY cạnh, không giấu
    ws.cell(r0, len(vps) + 2, "Ghi chú / lý do 未実施").font = Font(bold=True, color="FFFFFF")
    ws.cell(r0, len(vps) + 2).fill = PatternFill("solid", fgColor="2C3E50")

    # parity tóm tắt
    pr = r0 + len(engines) + 2
    ws.cell(pr, 1, "PARITY (so affordance giữa các engine cùng thời điểm)").font = Font(bold=True)
    for k, (vp, info) in enumerate(sorted((data.get("parity") or {}).items()), 1):
        ws.cell(pr + k, 1, vp)
        ws.cell(pr + k, 2, info.get("summary", "")).alignment = top
    for col, w in zip("ABCDEFG", (18, 16, 16, 16, 16, 52, 20)):
        ws.column_dimensions[col].width = w

    # ── Sheet 2: Findings ──
    fs = wb.create_sheet("Findings")
    for c, h in enumerate(HEADERS, 1):
        cell = fs.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50")
    r = 2
    allf = [(s, f) for s in data.get("screens", []) for f in s.get("findings", [])]
    for scr, fdg in sorted(allf, key=lambda x: SEV_ORDER.get(x[1].get("severity"), 9)):
        fs.cell(r, 1, scr.get("name", "")); fs.cell(r, 2, fdg.get("family", ""))
        fs.cell(r, 3, fdg.get("browser", "")); fs.cell(r, 4, fdg.get("viewport", ""))
        mc = fs.cell(r, 5, fdg.get("severity", ""))
        mc.fill = PatternFill("solid", fgColor=SEV_FILL.get(fdg.get("severity"), "BDC3C7"))
        fs.cell(r, 6, fdg.get("where", "")).alignment = top
        fs.cell(r, 7, fdg.get("observed", "")).alignment = top
        fs.cell(r, 8, fdg.get("fix", "")).alignment = top
        shot = fdg.get("shot", "")
        fs.cell(r, 9, os.path.basename(shot) if shot else "").alignment = top
        if shot and os.path.exists(os.path.join(base, shot)):
            img = XLImage(os.path.join(base, shot)); img.width = 240; img.height = 150
            fs.add_image(img, f"J{r}")
            fs.row_dimensions[r].height = 150
        r += 1
    if r == 2:
        fs.cell(2, 1, "Không có finding compatibility ở các engine ĐÃ chạy. "
                      "⚠️ Xem sheet 'Ma trận' để biết engine nào 未実施 — 未実施 KHÔNG phải PASS.")
    for col, w in zip("ABCDEFGHIJ", (20, 14, 11, 13, 8, 24, 46, 40, 26, 34)):
        fs.column_dimensions[col].width = w

    wb.save(out_path)
    n_un = sum(1 for c in cells if c.get("result") == "未実施")
    print(f"✅ compat report: {out_path} ({r - 2} finding · {len(cells)} ô · {n_un} ô 未実施)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_compat_report.py <results.json> <out.xlsx>")
    build(sys.argv[1], sys.argv[2])
