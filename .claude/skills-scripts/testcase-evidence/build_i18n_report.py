#!/usr/bin/env python3
"""build_i18n_report.py — i18n.results.json -> <Case>.i18n.xlsx (Ngôn ngữ + Findings).
Oracle = bất biến ngôn ngữ phổ quát (key-leak / mojibake / parity / locale-format). CƠ KHÍ, không reasoning.

Sheet "Ngôn ngữ" = bảng màn × probe, tô màu verdict. Sheet "Findings" = chi tiết lỗi + cách sửa.
localeFormat hiện cột riêng là WARN (không kéo verdict) — ghi rõ để người đọc không nhầm là FAIL.
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

V_FILL = {"PASS": "27AE60", "FAIL": "C0392B", "未実施": "7F8C8D"}
SEV_FILL = {"High": "C0392B", "Medium": "E67E22", "Low": "F1C40F"}
PROBE_COLS = [
    ("keyLeak", "Key thô", "hasLeak"),
    ("mojibake", "Mojibake", "hasMojibake"),
    ("parity", "Chưa dịch", None),
    ("localeFormat", "Format (WARN)", "warn"),
]

CAVEAT = ("Oracle = bất biến ngôn ngữ phổ quát (không cần SPEC, không cần ai đặt số). "
          "Data đã được MASK trước khi soi (tên viện/khách/mã = chữ Nhật hợp lệ, không tính là chưa-dịch). "
          "'Chưa dịch' chỉ đo được khi màn có ≥2 locale; app 1 ngôn ngữ → cột đó = 未実施. "
          "'Format (WARN)' là cảnh báo, KHÔNG kéo verdict xuống FAIL.")


def _probe_cell(ws, r, c, probe_key, probe, ctr):
    """Ô cho 1 probe: text + màu theo có-lỗi/không/không-áp-dụng."""
    if not probe:
        cell = ws.cell(r, c, "—"); cell.alignment = ctr; return
    if probe_key == "parity":
        if probe.get("inconclusive"):
            cell = ws.cell(r, c, "未実施"); cell.alignment = ctr
            cell.fill = PatternFill("solid", fgColor="7F8C8D"); cell.font = Font(color="FFFFFF")
            return
        n = len(probe.get("untranslated", []))
        bad = n > 0
        cell = ws.cell(r, c, f"{n} chưa dịch" if bad else "OK")
    elif probe_key == "localeFormat":
        bad = probe.get("warn")
        cell = ws.cell(r, c, "⚠ " + str(len(probe.get("hits", []))) if bad else "OK")
    else:
        flag = PROBE_COLS  # noqa
        bad = probe.get("hasLeak") or probe.get("hasMojibake")
        cell = ws.cell(r, c, str(len(probe.get("hits", []))) if bad else "OK")
    cell.alignment = ctr
    if probe_key == "localeFormat":
        if bad:
            cell.fill = PatternFill("solid", fgColor="F1C40F")
    else:
        cell.fill = PatternFill("solid", fgColor=("C0392B" if bad else "27AE60"))
        cell.font = Font(bold=True, color="FFFFFF")


def build(results_path, out_path):
    with open(results_path, encoding="utf-8") as f:
        data = json.load(f)
    top = Alignment(wrap_text=True, vertical="top")
    ctr = Alignment(horizontal="center", vertical="center")
    meta = data.get("meta", {})
    screens = data.get("screens", [])

    wb = Workbook(); ws = wb.active; ws.title = "Ngôn ngữ"
    ws.cell(1, 1, f"i18n / LOCALIZATION · {meta.get('case','')} · {meta.get('date','')} · {meta.get('tester','')}").font = Font(bold=True)
    c = ws.cell(2, 1, CAVEAT); c.alignment = top; c.font = Font(italic=True, size=9, color="8E44AD")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    ws.row_dimensions[2].height = 56

    hdr = ["Màn", "App", "Locale", "Verdict"] + [c[1] for c in PROBE_COLS]
    r0 = 4
    for j, h in enumerate(hdr, 1):
        cell = ws.cell(r0, j, h); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50"); cell.alignment = ctr

    rr = r0 + 1
    for s in screens:
        ws.cell(rr, 1, s.get("name", "")).alignment = top
        ws.cell(rr, 2, s.get("app", ""))
        ws.cell(rr, 3, s.get("locale", "")).alignment = ctr
        vc = ws.cell(rr, 4, s.get("result", "")); vc.alignment = ctr
        if s.get("result") in V_FILL:
            vc.fill = PatternFill("solid", fgColor=V_FILL[s["result"]]); vc.font = Font(bold=True, color="FFFFFF")
        probes = s.get("probes", {})
        for k, (pkey, _label, _flag) in enumerate(PROBE_COLS):
            _probe_cell(ws, rr, 5 + k, pkey, probes.get(pkey), ctr)
        rr += 1

    ws.cell(rr + 1, 1, "Chú giải: 🔴 = có lỗi (số = số chỗ) · 🟢 OK · ⚪ 未実施 = không đo được (vd app 1 ngôn ngữ) · "
                       "🟡 Format = WARN (không tính FAIL).").font = Font(italic=True, size=9)
    for col, w in zip("ABCDEFGH", (30, 12, 8, 10, 12, 12, 12, 14)):
        ws.column_dimensions[col].width = w

    # ── Sheet 2: Findings ──
    fs = wb.create_sheet("Findings")
    H = ["Màn", "Loại", "Mức", "Chi tiết (quan sát được)", "Cách sửa"]
    for j, h in enumerate(H, 1):
        cell = fs.cell(1, j, h); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50")
    r = 2
    for s in screens:
        for f in s.get("findings", []):
            fs.cell(r, 1, s.get("name", "")).alignment = top
            fs.cell(r, 2, f.get("family", "")).alignment = top
            mc = fs.cell(r, 3, f.get("severity", "")); mc.alignment = Alignment(horizontal="center")
            mc.fill = PatternFill("solid", fgColor=SEV_FILL.get(f.get("severity"), "BDC3C7"))
            fs.cell(r, 4, f.get("observed", "")).alignment = top
            fs.cell(r, 5, f.get("fix", "")).alignment = top
            fs.row_dimensions[r].height = 44
            r += 1
    if r == 2:
        fs.cell(2, 1, "Không tìm thấy lỗi ngôn ngữ ở các màn đã đo (trong phạm vi probe đã chạy).")
    for col, w in zip("ABCDE", (28, 14, 8, 70, 46)):
        fs.column_dimensions[col].width = w

    wb.save(out_path)
    nf = sum(len(s.get("findings", [])) for s in screens)
    print(f"✅ i18n report: {out_path} ({len(screens)} màn · {nf} finding)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_i18n_report.py <results.json> <out.xlsx>")
    build(sys.argv[1], sys.argv[2])
