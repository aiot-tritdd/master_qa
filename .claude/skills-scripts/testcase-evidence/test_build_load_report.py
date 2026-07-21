# Chạy: python3 test_build_load_report.py
import json, os, subprocess, sys, tempfile
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "build_load_report.py")


def main():
    results = {
        "meta": {"case": "T", "date": "2026-07-21", "tester": "QA", "mode": "stress",
                 "target": "reservation-dev", "endpoints": ["/api/a", "/api/b"]},
        "result": "FAIL",
        "metrics": {"errorRate": 0.03, "p95": 1900, "p99": 4000, "throughput": 120.5, "vusMax": 200},
        "thresholds": {"kind": "api", "p95": 800, "errorRate": 0.01},
        "stages": [
            {"vu": 50, "p95": 300, "errorRate": 0, "throughput": 90},
            {"vu": 100, "p95": 620, "errorRate": 0, "throughput": 110},
            {"vu": 150, "p95": 1900, "errorRate": 0.03, "throughput": 120},
        ],
        "knee": {"vu": 150, "reasons": ["p95 1900ms > 800ms"]},
        "bottleneck": {"hypothesis": "pool", "reason": "p95 nhảy bậc thang", "clientSideOnly": True},
        "caveats": ["client-side only", "DEV≠PROD"],
    }
    d = tempfile.mkdtemp()
    rp = os.path.join(d, "load.results.json")
    with open(rp, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False)
    out = os.path.join(d, "T.load.xlsx")
    subprocess.run([sys.executable, SCRIPT, rp, out], check=True)
    wb = load_workbook(out)
    assert "Kết quả" in wb.sheetnames and "Ramp" in wb.sheetnames, wb.sheetnames
    ws = wb["Kết quả"]
    assert ws.cell(4, 2).value == "FAIL", ws.cell(4, 2).value
    rs = wb["Ramp"]
    # dòng knee (vu=150) là dòng thứ 4 (header=1, 3 stage → rows 2,3,4)
    assert rs.cell(4, 1).value == 150, rs.cell(4, 1).value
    # có chữ "ĐIỂM GÃY" đâu đó
    found = any("ĐIỂM GÃY" in str(c.value) for row in rs.iter_rows() for c in row if c.value)
    assert found, "phải có dòng ĐIỂM GÃY"
    print("✅ test_build_load_report PASS")


if __name__ == "__main__":
    main()
