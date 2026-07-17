# Chạy: python3 test_build_security_report.py
import json, os, subprocess, sys, tempfile
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "build_security_report.py")

def _run(results, name="T.security.xlsx"):
    d = tempfile.mkdtemp()
    rp = os.path.join(d, "security.results.json")
    with open(rp, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False)
    out = os.path.join(d, name)
    subprocess.run([sys.executable, SCRIPT, rp, out], check=True)
    return load_workbook(out)

def test_findings_columns_and_split():
    results = {"meta": {"case": "TC11", "date": "2026-07-17", "tester": "QA-sec"},
        "screens": [{"name": "② 登録", "app": "ticket", "url": "/coupons/new/", "result": "FAIL",
            "findings": [
                {"family": "injection", "payloadClass": "sqli", "where": "ô Tên", "url": "/coupons/new/",
                 "severity": "Medium", "observed": "500 khi nhập '; DROP", "fix": "dùng prepared statement",
                 "shot": "shots/sec-TC11__coupon-new__sqli.png"},
                {"family": "xss", "payloadClass": "xss", "where": "ô Ghi chú", "url": "/coupons/new/",
                 "severity": "High", "observed": "payload execute (window.__SEC_XSS=1)", "fix": "escape output",
                 "shot": ""},
            ]}]}
    ws = _run(results)["Findings"]
    H = [ws.cell(1, c).value for c in range(1, 12)]
    assert H[3] == "Họ", H
    assert "Chi tiết" in H[7]
    assert H[8] == "Cách fix"
    assert H[9] == "File ảnh"
    # High phải sắp TRƯỚC Medium (severity-sort)
    assert ws.cell(2, 6).value == "High", ws.cell(2, 6).value
    # Chi tiết ≠ Cách fix (tách riêng)
    assert ws.cell(2, 8).value and ws.cell(2, 9).value and ws.cell(2, 8).value != ws.cell(2, 9).value
    # File ảnh = basename (dò ngược), rỗng khi không shot
    assert ws.cell(2, 10).value in ("", None)                # High finding không shot
    assert ws.cell(3, 10).value == "sec-TC11__coupon-new__sqli.png"

def test_coverage_sheet_and_empty():
    ws = _run({"meta": {"case": "E"}, "screens": [
        {"name": "① 一覧", "app": "ticket", "url": "/coupons/", "result": "PASS", "findings": []}]})
    assert "Đã quét" in ws.sheetnames
    cov = ws["Đã quét"]
    assert cov.cell(4, 4).value == "PASS"
    fw = ws["Findings"]
    assert fw.cell(2, 1).value == "Không có finding security nào ở các màn đã quét."

def main():
    test_findings_columns_and_split()
    test_coverage_sheet_and_empty()
    print("✅ test_build_security_report PASS (2 tests)")

if __name__ == "__main__":
    main()
