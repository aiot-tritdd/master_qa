# Chạy: python3 test_build_a11y_report.py
import json, os, subprocess, sys, tempfile
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "build_a11y_report.py")

def _run(results, name="T.a11y.xlsx"):
    d = tempfile.mkdtemp()
    rp = os.path.join(d, "a11y.results.json")
    with open(rp, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False)
    out = os.path.join(d, name)
    subprocess.run([sys.executable, SCRIPT, rp, out], check=True)
    return load_workbook(out)["Accessibility"]

def test_violations_rendered_and_sorted():
    results = {
        "meta": {"case": "T", "date": "2026-07-16", "tester": "QA-Server-a11y"},
        "screens": [{
            "name": "Đặt lịch", "app": "reservation", "url": "/booking", "result": "FAIL",
            "violations": [
                {"rule": "button-name", "impact": "serious", "wcag": "4.1.2",
                 "help": "Buttons must have discernible text",
                 "nodes": [{"target": "button.v-btn", "html": "<button></button>"}]},
                {"rule": "image-alt", "impact": "critical", "wcag": "1.1.1",
                 "help": "Images must have alternate text",
                 "description": "Ensure <img> elements have alternate text",
                 "helpUrl": "https://dequeuniversity.com/rules/axe/4.x/image-alt",
                 "nodes": [{"target": "img.a", "html": "<img>",
                            "failureSummary": "Fix any of the following:\n  Element has no alt attribute"},
                           {"target": "img.b", "html": "<img>"}]},
            ]
        }]
    }
    ws = _run(results)
    # layout cột: mã → WCAG → mức → số chỗ → lỗi gì → CHI TIẾT → CÁCH FIX (tách) → tài liệu → file ảnh → evidence
    assert ws.cell(1, 4).value == "Mã lỗi (rule)", ws.cell(1, 4).value
    assert ws.cell(1, 9).value == "Chi tiết (axe đo được)"
    assert ws.cell(1, 10).value == "Cách fix"
    assert ws.cell(1, 12).value == "File ảnh"
    # critical phải sắp TRƯỚC serious (impact-sort)
    assert ws.cell(2, 4).value == "image-alt", ws.cell(2, 4).value
    assert ws.cell(2, 5).value == "1.1.1"                             # WCAG (cột 5)
    assert ws.cell(2, 6).value == "critical"                         # Mức (cột 6)
    assert ws.cell(2, 7).value == 2                                   # Số chỗ = len(nodes)
    assert ws.cell(2, 8).value == "Ensure <img> elements have alternate text"   # Lỗi gì = description
    # image-alt CÓ trong RULE_GUIDE → Chi tiết + Cách fix TÁCH RIÊNG, khác nhau, không phải cùng 1 chuỗi
    detail, fix = ws.cell(2, 9).value, ws.cell(2, 10).value
    assert "alt" in detail.lower(), detail                            # chi tiết: thiếu alt
    assert "alt" in fix.lower() and detail != fix, (detail, fix)      # cách fix: thêm alt — KHÁC cột chi tiết
    assert ws.cell(2, 11).value == "https://dequeuniversity.com/rules/axe/4.x/image-alt"  # tài liệu (cột 11)
    assert ws.cell(2, 12).value in ("", None)                        # file ảnh trống khi không có shot
    # button-name cũng trong guide → có detail + fix riêng
    assert ws.cell(3, 4).value == "button-name"
    assert ws.cell(3, 6).value == "serious"
    assert ws.cell(3, 9).value and ws.cell(3, 10).value               # có cả chi tiết lẫn cách fix

def test_shot_filename_shown_for_traceability():
    results = {"meta": {"case": "TC11"}, "screens": [{
        "name": "① クーポン一覧", "app": "ticket", "url": "/coupons/", "result": "FAIL",
        "violations": [{"rule": "select-name", "impact": "critical", "wcag": "4.1.2", "help": "h",
                        "shot": "shots/a11y-TC11__coupon-list__select-name.png",
                        "nodes": [{"target": "select", "html": "<select>"}]}]}]}
    ws = _run(results)
    # File ảnh = basename (dò ngược Excel → shots/), KHÔNG phải path đầy đủ
    assert ws.cell(2, 12).value == "a11y-TC11__coupon-list__select-name.png", ws.cell(2, 12).value

def test_empty_screens_message():
    ws = _run({"meta": {"case": "E"}, "screens": []})
    assert ws.cell(2, 1).value == "Không có vi phạm a11y nào ở các màn đã quét."

def test_coverage_sheet():
    results = {"meta": {"case": "T", "date": "2026-07-16", "tester": "QA"},
        "screens": [
            {"name": "Đặt lịch", "app": "reservation", "url": "/booking", "result": "FAIL",
             "violations": [{"rule": "image-alt", "impact": "critical", "wcag": "1.1.1", "help": "h",
                             "nodes": [{"target": "img", "html": "<img>"}]}]},
            {"name": "Login", "app": "pro", "url": "/login", "result": "PASS", "violations": []},
        ]}
    d = tempfile.mkdtemp()
    rp = os.path.join(d, "r.json")
    with open(rp, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False)
    out = os.path.join(d, "c.xlsx")
    subprocess.run([sys.executable, SCRIPT, rp, out], check=True)
    from openpyxl import load_workbook as _lw
    wb = _lw(out)
    assert "Màn quét" in wb.sheetnames, wb.sheetnames
    sw = wb["Màn quét"]
    rows = {sw.cell(r, 1).value: (sw.cell(r, 4).value, sw.cell(r, 5).value) for r in range(4, 6)}
    assert rows.get("Login") == ("PASS", 0), rows
    assert rows.get("Đặt lịch")[0] == "FAIL", rows

def main():
    test_violations_rendered_and_sorted()
    test_shot_filename_shown_for_traceability()
    test_empty_screens_message()
    test_coverage_sheet()
    print("✅ test_build_a11y_report PASS (4 tests)")

if __name__ == "__main__":
    main()
