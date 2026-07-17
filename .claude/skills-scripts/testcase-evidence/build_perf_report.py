#!/usr/bin/env python3
"""build_perf_report.py — perf.results.json -> <Case>.perf.xlsx (Chỉ số + Findings).
Oracle = Core Web Vitals (ngưỡng Google công bố). CƠ KHÍ, không reasoning.

Sheet "Chỉ số" là bảng màn × metric, tô màu theo good/needs-improvement/poor, KÈM ngưỡng ngay
trong bảng để người đọc không phải tra đâu khác — và kèm cảnh báo LAB≠FIELD ngay đầu sheet.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

RATE_FILL = {"good": "27AE60", "needs-improvement": "E67E22", "poor": "C0392B", "unknown": "7F8C8D"}
SEV_FILL = {"High": "C0392B", "Medium": "E67E22", "Low": "F1C40F"}
V_FILL = {"PASS": "27AE60", "FAIL": "C0392B", "未実施": "7F8C8D"}
ORDER = ["LCP", "CLS", "TBT", "FCP", "TTFB"]

CAVEAT = ("⚠️ ĐÂY LÀ ĐO LAB (1 máy, 1 đường mạng, trên DEV) — KHÔNG phải đo người dùng thật. "
          "Chuẩn Core Web Vitals thật là phân vị 75 của người dùng thật (CrUX): điện thoại yếu, 4G. "
          "⇒ Lab \"good\" KHÔNG chứng minh người dùng thật thấy nhanh; nhưng lab \"poor\" thì CHẮC CHẮN có vấn đề. "
          "Kết quả này là CẬN DƯỚI của mức tệ. Dev ≠ prod (không CDN, data ít). "
          "INP không đo được ở lab → dùng TBT làm proxy (khuyến nghị của Google).")


def build(results_path, out_path):
    with open(results_path, encoding="utf-8") as f:
        data = json.load(f)
    top = Alignment(wrap_text=True, vertical="top")
    ctr = Alignment(horizontal="center", vertical="center")
    meta = data.get("meta", {})
    screens = data.get("screens", [])

    wb = Workbook(); ws = wb.active; ws.title = "Chỉ số"
    ws.cell(1, 1, f"PERFORMANCE · {meta.get('case','')} · {meta.get('date','')} · {meta.get('tester','')}").font = Font(bold=True)
    c = ws.cell(2, 1, CAVEAT); c.alignment = top; c.font = Font(italic=True, size=9, color="C0392B")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    ws.row_dimensions[2].height = 58

    hdr = ["Màn", "App", "Verdict"] + ORDER
    r0 = 4
    for j, h in enumerate(hdr, 1):
        cell = ws.cell(r0, j, h); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50"); cell.alignment = ctr

    # dòng ngưỡng — để ngay dưới header, khỏi phải tra chỗ khác
    ws.cell(r0 + 1, 1, "Ngưỡng Google (good ≤)").font = Font(italic=True, bold=True)
    TH = data.get("thresholds", {})
    for j, m in enumerate(ORDER, 4):
        t = TH.get(m, {})
        cc = ws.cell(r0 + 1, j, f"≤ {t.get('good','?')}{t.get('unit','')}")
        cc.alignment = ctr; cc.font = Font(italic=True)

    rr = r0 + 2
    for s in screens:
        ws.cell(rr, 1, s.get("name", "")).alignment = top
        ws.cell(rr, 2, s.get("app", ""))
        vc = ws.cell(rr, 3, s.get("result", "")); vc.alignment = ctr
        if s.get("result") in V_FILL:
            vc.fill = PatternFill("solid", fgColor=V_FILL[s["result"]]); vc.font = Font(bold=True, color="FFFFFF")
        agg = s.get("metrics", {})
        for j, m in enumerate(ORDER, 4):
            d = agg.get(m) or {}
            med = d.get("median")
            txt = "—" if med is None else (f"{med:.3f}" if m == "CLS" else f"{round(med)}{TH.get(m,{}).get('unit','')}")
            if d.get("unstable"):
                txt += " ⚠️nhiễu"
            cc = ws.cell(rr, j, txt); cc.alignment = ctr
            rt = d.get("rating", "unknown")
            if rt in RATE_FILL and med is not None:
                cc.fill = PatternFill("solid", fgColor=RATE_FILL[rt])
                cc.font = Font(bold=True, color="FFFFFF")
        rr += 1

    ws.cell(rr + 1, 1, "Chú giải: 🟢 good = đạt chuẩn Google · 🟠 needs-improvement = KHÔNG đạt · 🔴 poor = tệ. "
                       "⚠️nhiễu = dao động giữa các lần chạy lớn hơn trung vị → đừng tin con số chính xác.").font = Font(italic=True, size=9)
    for col, w in zip("ABCDEFGH", (30, 12, 10, 14, 14, 14, 14, 14)):
        ws.column_dimensions[col].width = w

    # ── Sheet 2: Findings ──
    fs = wb.create_sheet("Findings")
    H = ["Màn", "Chỉ số", "Mức", "Chi tiết (số đo + ngưỡng + dao động)", "Chuẩn & cách sửa"]
    for j, h in enumerate(H, 1):
        cell = fs.cell(1, j, h); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50")
    r = 2
    for s in screens:
        for f in s.get("findings", []):
            fs.cell(r, 1, s.get("name", "")).alignment = top
            fs.cell(r, 2, f.get("metric", "")).alignment = ctr
            mc = fs.cell(r, 3, f.get("severity", "")); mc.alignment = ctr
            mc.fill = PatternFill("solid", fgColor=SEV_FILL.get(f.get("severity"), "BDC3C7"))
            fs.cell(r, 4, f.get("observed", "")).alignment = top
            fs.cell(r, 5, f.get("fix", "")).alignment = top
            fs.row_dimensions[r].height = 46
            r += 1
    if r == 2:
        fs.cell(2, 1, "Mọi chỉ số đều đạt chuẩn Google ở các màn đã đo. ⚠️ Nhớ: đây là đo LAB trên dev — xem cảnh báo sheet 'Chỉ số'.")
    for col, w in zip("ABCDE", (28, 10, 8, 72, 52)):
        fs.column_dimensions[col].width = w

    wb.save(out_path)
    nf = sum(len(s.get("findings", [])) for s in screens)
    print(f"✅ perf report: {out_path} ({len(screens)} màn · {nf} finding)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_perf_report.py <results.json> <out.xlsx>")
    build(sys.argv[1], sys.argv[2])
