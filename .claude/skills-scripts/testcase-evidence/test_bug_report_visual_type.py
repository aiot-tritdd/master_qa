# Chạy: python3 test_bug_report_visual_type.py
import copy, json, os, subprocess, sys, tempfile
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
REAL = os.path.join(HERE, "..", "..", "..", "wtf-is-this", "bug-he-thong.tcs.json")

def main():
    with open(REAL, encoding="utf-8") as f:
        data = json.load(f)
    a = copy.deepcopy(data["tcs"][0])
    a.update({"bug_id": "BUG-VIS-TEST", "bug_type": "Visual", "result": "FAIL",
              "screen": "Pro — Đặt lịch", "status": "Mở", "found_at": "2026-07-16"})
    data["tcs"].append(a)
    d = tempfile.mkdtemp()
    jp = os.path.join(d, "bugs.tcs.json")
    with open(jp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    out = os.path.join(d, "bugs.xlsx")
    subprocess.run([sys.executable, os.path.join(HERE, "build_bug_report.py"), jp, out], check=True)
    wb = load_workbook(out); ws = wb[wb.sheetnames[0]]
    found = any(c.value == "Visual" for row in ws.iter_rows() for c in row)
    assert found, "Tổng quan phải hiện nhãn 'Visual'"
    print("✅ test_bug_report_visual_type PASS")

if __name__ == "__main__":
    main()
