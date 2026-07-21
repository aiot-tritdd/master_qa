#!/usr/bin/env python3
"""build_load_report.py — load.results.json -> <Case>.load.xlsx (Kết quả + Ramp).
GREY-BOX: chỉ đo CLIENT-SIDE (k6). Không thấy ruột server (không đụng AWS khách). CƠ KHÍ, không reasoning.

Sheet "Kết quả": verdict + số đo vs ngưỡng + KHỐI CẢNH BÁO to (client-side only / DEV≠PROD / p95 placeholder).
Sheet "Ramp" (nếu có stages): bảng VU × chỉ số, tô đỏ mức GÃY, kèm GIẢ THUYẾT nút thắt (không phải kết luận).
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

V_FILL = {"PASS": "27AE60", "FAIL": "C0392B", "未実施": "7F8C8D"}

BANNER = ("⚠️ GREY-BOX · CHỈ ĐO CLIENT-SIDE. Đây là số nhìn từ ngoài (k6): latency/error/throughput. "
          "KHÔNG thấy ruột server (CPU/RAM/DB pool) vì không đụng AWS khách ⇒ phần 'nút thắt' bên dưới chỉ là "
          "GIẢ THUYẾT suy từ hình đường cong, KHÔNG phải kết luận. DEV ≠ PROD (không CDN, data ít, 1 instance). "
          "Ngưỡng p95 là PLACEHOLDER mượn từ CWV — cần SLA business (đỉnh concurrency + p95 mục tiêu) mới chốt được.")

BN_HYP = {"cpu": "CPU/worker (nhìn EC2 CPU mới chắc)", "pool": "cạn connection pool (nhìn RDS DatabaseConnections mới chắc)",
          "worker": "worker chết/quá tải (nhìn EC2 mới chắc)", "db-lock": "chờ khoá DB (nhìn RDS mới chắc)",
          "inconclusive": "chưa rõ"}


def build(results_path, out_path):
    with open(results_path, encoding="utf-8") as f:
        data = json.load(f)
    top = Alignment(wrap_text=True, vertical="top")
    ctr = Alignment(horizontal="center", vertical="center")
    meta = data.get("meta", {})
    mode = meta.get("mode", "load")

    wb = Workbook(); ws = wb.active; ws.title = "Kết quả"
    ws.cell(1, 1, f"{mode.upper()} TEST · {meta.get('case','')} · {meta.get('date','')} · {meta.get('tester','')} "
                  f"· target: {meta.get('target','')} · endpoints: {', '.join(meta.get('endpoints', []))}").font = Font(bold=True)
    c = ws.cell(2, 1, BANNER); c.alignment = top; c.font = Font(italic=True, size=9, color="C0392B")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    ws.row_dimensions[2].height = 74

    vc = ws.cell(4, 1, "Verdict"); vc.font = Font(bold=True)
    rc = ws.cell(4, 2, data.get("result", "")); rc.alignment = ctr
    if data.get("result") in V_FILL:
        rc.fill = PatternFill("solid", fgColor=V_FILL[data["result"]]); rc.font = Font(bold=True, color="FFFFFF")

    th = data.get("thresholds", {})
    m = data.get("metrics", {})
    rows = [
        ("error rate", f'{(m.get("errorRate") or 0)*100:.2f}%' if m.get("errorRate") is not None else "—",
         f'< {(th.get("errorRate") or 0.01)*100:.0f}%'),
        ("p95 latency", f'{round(m["p95"])}ms' if m.get("p95") is not None else "—",
         f'< {th.get("p95","?")}ms ({th.get("kind","api")}) — placeholder'),
        ("p99 latency", f'{round(m["p99"])}ms' if m.get("p99") is not None else "—", ""),
        ("throughput", f'{m.get("throughput"):.1f} req/s' if m.get("throughput") is not None else "—", ""),
        ("VU tối đa", str(m.get("vusMax", "—")), ""),
    ]
    r0 = 6
    for j, h in enumerate(("Chỉ số", "Đo được", "Ngưỡng"), 1):
        cell = ws.cell(r0, j, h); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50"); cell.alignment = ctr
    for i, (k, v, t) in enumerate(rows, 1):
        ws.cell(r0 + i, 1, k); ws.cell(r0 + i, 2, v).alignment = ctr
        ws.cell(r0 + i, 3, t).alignment = top
    for col, w in zip("ABC", (18, 20, 40)):
        ws.column_dimensions[col].width = w

    # ── Sheet 2: Ramp (stress) ──
    stages = data.get("stages", [])
    if stages:
        rs = wb.create_sheet("Ramp")
        for j, h in enumerate(("VU", "p95 (ms)", "error %", "throughput (req/s)"), 1):
            cell = rs.cell(1, j, h); cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2C3E50"); cell.alignment = ctr
        knee = data.get("knee") or {}
        knee_vu = knee.get("vu")
        for i, st in enumerate(stages, 2):
            rs.cell(i, 1, st.get("vu")).alignment = ctr
            rs.cell(i, 2, round(st["p95"]) if st.get("p95") is not None else "—").alignment = ctr
            rs.cell(i, 3, f'{(st.get("errorRate") or 0)*100:.2f}').alignment = ctr
            rs.cell(i, 4, round(st.get("throughput"), 1) if st.get("throughput") is not None else "—").alignment = ctr
            if st.get("vu") == knee_vu:
                for j in range(1, 5):
                    rs.cell(i, j).fill = PatternFill("solid", fgColor="C0392B"); rs.cell(i, j).font = Font(bold=True, color="FFFFFF")
        r = len(stages) + 3
        if knee_vu is not None:
            rs.cell(r, 1, f"ĐIỂM GÃY: {knee_vu} VU — {'; '.join(knee.get('reasons', []))}").font = Font(bold=True, color="C0392B")
        else:
            rs.cell(r, 1, "Không gãy trong dải VU đã bắn.").font = Font(bold=True, color="27AE60")
        bo = data.get("bottleneck", {})
        rs.cell(r + 1, 1, f"GIẢ THUYẾT nút thắt: {bo.get('hypothesis','?')} → {BN_HYP.get(bo.get('hypothesis'),'')}. "
                          f"{bo.get('reason','')} (client-side only — CHƯA xác nhận bằng metrics server).").alignment = top
        rs.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=4)
        rs.row_dimensions[r + 1].height = 42
        for col, w in zip("ABCD", (10, 12, 10, 20)):
            rs.column_dimensions[col].width = w

    wb.save(out_path)
    print(f"✅ {mode} report: {out_path} (verdict {data.get('result','?')} · {len(stages)} stage)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_load_report.py <results.json> <out.xlsx>")
    build(sys.argv[1], sys.argv[2])
