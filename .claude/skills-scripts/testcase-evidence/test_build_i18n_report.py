# Chạy: python3 test_build_i18n_report.py
import json, os, subprocess, sys, tempfile
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "build_i18n_report.py")


def main():
    results = {"meta": {"case": "T", "date": "2026-07-21", "tester": "QA"},
        "screens": [
            {"name": "Widget EN", "app": "reservation", "url": "/en/2", "locale": "en", "result": "FAIL",
             "probes": {
                 "keyLeak": {"ran": True, "hasLeak": True, "hits": [{"text": "reservation.selectCourse"}]},
                 "mojibake": {"ran": True, "hasMojibake": False, "hits": []},
                 "parity": {"inconclusive": False, "untranslated": ["コース選択"]},
                 "localeFormat": {"warn": True, "hits": [{"text": "1,000円"}]},
             },
             "findings": [{"family": "key-leak", "severity": "High",
                           "observed": "Nút hiện key thô 'reservation.selectCourse'", "fix": "Bổ sung bản dịch"}]},
            {"name": "Pro JA", "app": "pro", "url": "/reservations", "locale": "ja", "result": "PASS",
             "probes": {
                 "keyLeak": {"ran": True, "hasLeak": False, "hits": []},
                 "mojibake": {"ran": True, "hasMojibake": False, "hits": []},
                 "parity": {"inconclusive": True, "untranslated": [], "reason": "app 1 ngôn ngữ"},
                 "localeFormat": {"warn": False, "hits": [], "skipped": True},
             },
             "findings": []},
        ]}
    d = tempfile.mkdtemp()
    rp = os.path.join(d, "i18n.results.json")
    with open(rp, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False)
    out = os.path.join(d, "T.i18n.xlsx")
    subprocess.run([sys.executable, SCRIPT, rp, out], check=True)
    wb = load_workbook(out)
    assert "Ngôn ngữ" in wb.sheetnames and "Findings" in wb.sheetnames, wb.sheetnames
    ws = wb["Ngôn ngữ"]
    assert ws.cell(4, 1).value == "Màn", ws.cell(4, 1).value
    assert ws.cell(5, 4).value == "FAIL", ws.cell(5, 4).value          # verdict màn 1
    assert ws.cell(6, 4).value == "PASS", ws.cell(6, 4).value          # verdict màn 2
    # màn Pro JA: cột parity (cột 7) phải là 未実施 (app 1 ngôn ngữ)
    assert ws.cell(6, 7).value == "未実施", ws.cell(6, 7).value
    fs = wb["Findings"]
    assert fs.cell(2, 2).value == "key-leak", fs.cell(2, 2).value
    print("✅ test_build_i18n_report PASS")


if __name__ == "__main__":
    main()
