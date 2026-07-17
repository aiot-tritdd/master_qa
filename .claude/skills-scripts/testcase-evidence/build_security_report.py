#!/usr/bin/env python3
"""build_security_report.py — security.results.json -> <Case>.security.xlsx (Findings + Đã quét).
Oracle = bất biến an ninh phổ quát (không phải spec). CƠ KHÍ, không reasoning."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage

SEV_ORDER = {"High": 0, "Medium": 1, "Low": 2}
SEV_FILL = {"High": "C0392B", "Medium": "E67E22", "Low": "F1C40F"}
HEADERS = ["Màn", "App", "URL", "Họ", "Payload-class", "Mức", "Nơi (field/endpoint)",
           "Chi tiết (quan sát được)", "Cách fix", "File ảnh", "Evidence"]


def build(results_path, out_path):
    with open(results_path, encoding="utf-8") as f:
        data = json.load(f)
    base = os.path.dirname(os.path.abspath(results_path))
    top = Alignment(wrap_text=True, vertical="top")
    wb = Workbook(); ws = wb.active; ws.title = "Findings"
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50")
    r = 2
    for scr in data.get("screens", []):
        for fdg in sorted(scr.get("findings", []), key=lambda x: SEV_ORDER.get(x.get("severity"), 9)):
            ws.cell(r, 1, scr.get("name", "")); ws.cell(r, 2, scr.get("app", "")); ws.cell(r, 3, scr.get("url", ""))
            ws.cell(r, 4, fdg.get("family", "")); ws.cell(r, 5, fdg.get("payloadClass", ""))
            mc = ws.cell(r, 6, fdg.get("severity", ""))
            mc.fill = PatternFill("solid", fgColor=SEV_FILL.get(fdg.get("severity"), "BDC3C7"))
            ws.cell(r, 7, fdg.get("where", "")).alignment = top
            ws.cell(r, 8, fdg.get("observed", "")).alignment = top   # Chi tiết (quan sát)
            ws.cell(r, 9, fdg.get("fix", "")).alignment = top        # Cách fix (tách riêng)
            shot = fdg.get("shot", "")
            ws.cell(r, 10, os.path.basename(shot) if shot else "").alignment = top
            if shot and os.path.exists(os.path.join(base, shot)):
                img = XLImage(os.path.join(base, shot)); img.width = 240; img.height = 150
                ws.add_image(img, f"K{r}")
            ws.row_dimensions[r].height = 150
            r += 1
    if r == 2:
        ws.cell(2, 1, "Không có finding security nào ở các màn đã quét.")
    for col, w in zip("ABCDEFGHIJK", (14, 9, 20, 14, 13, 8, 22, 40, 40, 30, 34)):
        ws.column_dimensions[col].width = w

    cov = wb.create_sheet("Đã quét")
    meta = data.get("meta", {})
    cov.cell(1, 1, f"SECURITY · {meta.get('case','')} · {meta.get('date','')} · {meta.get('tester','')}").font = Font(bold=True)
    for c, h in enumerate(["Màn", "App", "URL", "Verdict", "Finding High"], 1):
        cell = cov.cell(3, c, h); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50")
    rr = 4
    for scr in data.get("screens", []):
        hi = sum(1 for x in scr.get("findings", []) if x.get("severity") == "High")
        cov.cell(rr, 1, scr.get("name", "")); cov.cell(rr, 2, scr.get("app", "")); cov.cell(rr, 3, scr.get("url", ""))
        cov.cell(rr, 4, scr.get("result", "")); cov.cell(rr, 5, hi)
        rr += 1
    for col, w in zip("ABCDE", (18, 10, 24, 10, 14)):
        cov.column_dimensions[col].width = w

    wb.save(out_path)
    print(f"✅ security report: {out_path} ({r - 2} finding)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_security_report.py <results.json> <out.xlsx>")
    build(sys.argv[1], sys.argv[2])
