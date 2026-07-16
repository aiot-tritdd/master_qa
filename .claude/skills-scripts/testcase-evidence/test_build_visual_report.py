# Chạy: python3 test_build_visual_report.py
import json, os, subprocess, sys, tempfile
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "build_visual_report.py")

def main():
    results = {"meta": {"case": "T", "date": "2026-07-16", "tester": "QA"},
        "screens": [
            {"name": "Đặt lịch", "app": "pro", "url": "/reservations", "result": "FAIL",
             "diff_ratio": 0.032, "baseline": "", "current": "", "diff": ""},
            {"name": "Login", "app": "pro", "url": "/login", "result": "NEW-BASELINE",
             "diff_ratio": 0, "baseline": "", "current": "", "diff": ""},
        ]}
    d = tempfile.mkdtemp()
    rp = os.path.join(d, "visual.results.json")
    with open(rp, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False)
    out = os.path.join(d, "T.visual.xlsx")
    subprocess.run([sys.executable, SCRIPT, rp, out], check=True)
    wb = load_workbook(out)
    assert "Visual" in wb.sheetnames and "Màn quét" in wb.sheetnames, wb.sheetnames
    ws = wb["Visual"]
    assert ws.cell(1, 1).value == "Màn", ws.cell(1, 1).value
    assert ws.cell(2, 4).value == "FAIL", ws.cell(2, 4).value
    cov = wb["Màn quét"]
    rows = {cov.cell(r, 1).value: cov.cell(r, 4).value for r in range(4, 6)}
    assert rows.get("Login") == "NEW-BASELINE", rows
    print("✅ test_build_visual_report PASS")

if __name__ == "__main__":
    main()
