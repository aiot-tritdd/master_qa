#!/usr/bin/env python3
"""build_visual_report.py — visual.results.json -> <Case>.visual.xlsx (2 sheet: Visual + Màn quét).
Oracle = baseline người-duyệt. CƠ KHÍ, không reasoning."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage

HEAD = "2C3E50"
VERDICT_FILL = {"FAIL": "C0392B", "PASS": "27AE60", "NEW-BASELINE": "F1C40F", "未実施": "BDC3C7"}
HEADERS = ["Màn", "App", "URL", "Verdict", "Diff %", "Baseline", "Hiện tại", "Diff"]

def _img(ws, cell, base_dir, rel, r):
    if rel and os.path.exists(os.path.join(base_dir, rel)):
        im = XLImage(os.path.join(base_dir, rel)); im.width = 200; im.height = 130
        ws.add_image(im, cell); ws.row_dimensions[r].height = 105

def build(results_path, out_path):
    with open(results_path, encoding="utf-8") as f:
        data = json.load(f)
    base_dir = os.path.dirname(os.path.abspath(results_path))
    meta = data.get("meta", {}); screens = data.get("screens", [])
    wb = Workbook(); ws = wb.active; ws.title = "Visual"
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, c, h); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEAD)
    r = 2
    for s in screens:
        ws.cell(r, 1, s.get("name", "")); ws.cell(r, 2, s.get("app", "")); ws.cell(r, 3, s.get("url", ""))
        vc = ws.cell(r, 4, s.get("result", ""))
        vc.fill = PatternFill("solid", fgColor=VERDICT_FILL.get(s.get("result"), "BDC3C7"))
        ws.cell(r, 5, round(s.get("diff_ratio", 0) * 100, 3))
        _img(ws, f"F{r}", base_dir, s.get("baseline", ""), r)
        _img(ws, f"G{r}", base_dir, s.get("current", ""), r)
        _img(ws, f"H{r}", base_dir, s.get("diff", ""), r)
        r += 1
    if r == 2:
        ws.cell(2, 1, "Không có màn nào được quét.")
    for col, w in zip("ABCDEFGH", (18, 10, 24, 14, 9, 30, 30, 30)):
        ws.column_dimensions[col].width = w
    # Sheet coverage
    cov = wb.create_sheet("Màn quét")
    cov.cell(1, 1, f"VISUAL · {meta.get('case','')} · {meta.get('date','')} · {meta.get('tester','')}").font = Font(bold=True)
    for c, h in enumerate(["Màn", "App", "URL", "Verdict", "Diff %"], 1):
        cell = cov.cell(3, c, h); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEAD)
    rr = 4
    for s in screens:
        cov.cell(rr, 1, s.get("name", "")); cov.cell(rr, 2, s.get("app", "")); cov.cell(rr, 3, s.get("url", ""))
        cov.cell(rr, 4, s.get("result", "")); cov.cell(rr, 5, round(s.get("diff_ratio", 0) * 100, 3))
        rr += 1
    for col, w in zip("ABCDE", (18, 10, 24, 14, 9)):
        cov.column_dimensions[col].width = w
    wb.save(out_path)
    print(f"✅ visual report: {out_path} ({len(screens)} màn)")

if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_visual_report.py <results.json> <out.xlsx>")
    build(sys.argv[1], sys.argv[2])
