# Chạy: python3 test_bug_report_a11y_type.py
# Dùng dữ liệu THẬT làm nền (records đầy đủ field) + thêm 1 bug Accessibility.
import copy, json, os, subprocess, sys, tempfile
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
REAL = os.path.join(HERE, "..", "..", "..", "wtf-is-this", "bug-he-thong.tcs.json")

def main():
    with open(REAL, encoding="utf-8") as f:
        data = json.load(f)
    a = copy.deepcopy(data["tcs"][0])
    a.update({"bug_id": "BUG-A11Y-TEST", "bug_type": "Accessibility", "result": "FAIL",
              "screen": "Reservation — Đặt lịch", "status": "Mở", "found_at": "2026-07-16"})
    data["tcs"].append(a)
    with tempfile.TemporaryDirectory() as d:
        jp = os.path.join(d, "bugs.tcs.json")
        with open(jp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        out = os.path.join(d, "bugs.xlsx")
        subprocess.run([sys.executable, os.path.join(HERE, "build_bug_report.py"), jp, out], check=True)
        wb = load_workbook(out)
        ws = wb[wb.sheetnames[0]]  # 'Tổng quan'
        found = any(c.value == "Accessibility" for row in ws.iter_rows() for c in row)
        assert found, "Tổng quan phải hiện nhãn 'Accessibility' (đã thêm vào tuple type)"
    print("✅ test_bug_report_a11y_type PASS")

if __name__ == "__main__":
    main()
