# Chạy: python3 test_bug_report_security_type.py
# Dùng dữ liệu THẬT làm nền (records đầy đủ field → đủ data cho decor/chart) + thêm 1 bug Security.
import copy, json, os, subprocess, sys, tempfile
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
REAL = os.path.join(HERE, "..", "..", "..", "wtf-is-this", "bug-he-thong.tcs.json")

def main():
    with open(REAL, encoding="utf-8") as f:
        data = json.load(f)
    s = copy.deepcopy(data["tcs"][0])
    s.update({"bug_id": "BUG-SEC-TEST", "bug_type": "Security", "result": "FAIL",
              "screen": "Ticket — ② 登録", "status": "Mở", "found_at": "2026-07-17",
              "title": "XSS: payload execute ở ô Ghi chú"})
    data["tcs"].append(s)
    with tempfile.TemporaryDirectory() as d:
        jp = os.path.join(d, "bugs.tcs.json")
        with open(jp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        out = os.path.join(d, "bugs.xlsx")
        subprocess.run([sys.executable, os.path.join(HERE, "build_bug_report.py"), jp, out], check=True)
        wb = load_workbook(out)
        ws = wb[wb.sheetnames[0]]  # 'Tổng quan'
        found = any(c.value == "Security" for row in ws.iter_rows() for c in row)
        assert found, "Tổng quan phải hiện nhãn 'Security' (đã thêm vào tuple type)"
    print("✅ test_bug_report_security_type PASS")

if __name__ == "__main__":
    main()
