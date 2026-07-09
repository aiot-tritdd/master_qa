#!/usr/bin/env python3
"""build_evidence.py — tcs.json (+ shots/) -> .xlsx chuẩn TEMPLATE (3 sheet). CƠ KHÍ, KHÔNG reasoning.

Usage: python3 build_evidence.py <folder>/tcs.json <folder>/<Ten>.xlsx
Deps:  pip install openpyxl pillow

FORMAT = theme.json (cạnh file này) — nguồn sự thật DUY NHẤT cho màu/layout/nhãn.
Đổi giao diện file xuất => sửa theme.json, KHÔNG sửa file .py này.

Xuất 3 sheet:
  - Cover      : report header + SUMMARY (COUNTIF Total/PASS/FAIL/未実施/SPEC-GAP)
  - Test Cases : mỗi case = 1 BLOCK DỌC (Test Case/Pre/Steps/Expected/Actual/Note)
                 + Evidence (Before) [E:J] / (After) [K:P] nhúng ẢNH TO (PNG rõ).
  - Checklist  : # | TC ID | Title | Priority | Result | Nguồn (alt-row + Total formulas)

result ∈ {PASS, FAIL, 未実施, SPEC-GAP}. SPEC-GAP = spec không định nghĩa kỳ vọng ở chỗ
METHOD.md bảo phải có case → không bịa expect → không chấm được (KHÁC 未実施 = không quan sát được).

Đọc {meta, shots_dir, tcs:[{id,screen,pri,result,title,pre,steps,expect,actual,note,source,before,after}]}
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None

# ---------------- theme (single source of truth cho FORMAT) ----------------
THEME = json.loads((Path(__file__).parent / "theme.json").read_text(encoding="utf-8"))
P, FT, LY, LB = THEME["palette"], THEME["fonts"], THEME["layout"], THEME["labels"]

WHITE, NAVY, BLUE2, LABEL = P["white"], P["navy"], P["blue"], P["lblue"]
NUMF, NOTEF, NAF, ALT = P["num_bg"], P["note_bg"], P["na_bg"], P["alt_row"]

RES_FILL = {k: P[v["bg"]] for k, v in LB["status"].items()}
RES_FONT = {k: P[v["tx"]] for k, v in LB["status"].items()}
PRI = {k: P[v] for k, v in LB["priority"].items()}  # gồm cả "_default"

_TH = Side(style="thin", color=P["border"])
BORD = Border(left=_TH, right=_TH, top=_TH, bottom=_TH)

TOP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)

IMG_W = LY["image_display_width_px"]
EB0, EB1 = LY["evidence_before"]
EA0, EA1 = LY["evidence_after"]
LAST_COL = EA1
FIELDS = [tuple(f) for f in LB["fields"]]


def _f(hex_):
    return PatternFill("solid", fgColor=hex_)


def _c(ws, r, c, v=None, *, font=None, fill=None, align=TOP, border=True):
    x = ws.cell(row=r, column=c, value=v)
    if font:
        x.font = font
    if fill:
        x.fill = _f(fill)
    x.alignment = align
    if border:
        x.border = BORD
    return x


def _img_size(p):
    if PILImage is None:
        return IMG_W, int(IMG_W * 0.62)
    with PILImage.open(p) as im:
        ratio = (im.height / im.width) if im.width else 0.62
    return IMG_W, int(IMG_W * ratio)


# ---------------- Cover ----------------
def build_cover(wb, meta):
    ws = wb.active
    ws.title = "Cover"
    for col, w in LY["cover_cols"].items():
        ws.column_dimensions[col].width = w
    ws.merge_cells("B2:C3")
    _c(ws, 2, 2, LB["cover_title"], font=Font(bold=True, size=FT["cover_title"], color=WHITE),
       fill=NAVY, align=CEN, border=False)
    keys = ["project", "module", "issue", "tester", "date", "env"]
    r = 5
    for label, key in zip(LB["cover_rows"], keys):
        _c(ws, r, 2, label, font=Font(bold=True), fill=LABEL)
        _c(ws, r, 3, str(meta.get(key, "")))
        r += 1
    r += 1
    ws.merge_cells(f"B{r}:C{r}")
    _c(ws, r, 2, LB["cover_summary"], font=Font(bold=True, color=WHITE), fill=BLUE2, align=CEN)
    r += 1
    for label, match, col, fill_key in LB["cover_summary_rows"]:
        _c(ws, r, 2, label, font=Font(bold=True), fill=LABEL)
        _c(ws, r, 3, f"=COUNTIF('Test Cases'!{col}:{col},\"{match}\")",
           font=Font(bold=True), fill=P[fill_key])
        r += 1


# ---------------- Test Cases ----------------
def build_testcases(wb, meta, tcs, shots_dir):
    ws = wb.create_sheet("Test Cases")
    for col, w in LY["detail_cols"].items():
        ws.column_dimensions[col].width = w
    for i in range(EB0, LAST_COL + 1):
        ws.column_dimensions[get_column_letter(i)].width = LY["evidence_col_width"]

    end = get_column_letter(LAST_COL)
    ws.merge_cells(f"A1:{end}1")
    _c(ws, 1, 1, f"TEST CASES — {meta.get('module', '')}",
       font=Font(bold=True, size=FT["sheet_title"], color=WHITE), fill=NAVY,
       align=Alignment(vertical="center"))
    hdr = Font(bold=True, color=WHITE)
    _c(ws, 2, 1, "#", font=hdr, fill=NAVY, align=CEN)
    _c(ws, 2, 2, "Field", font=hdr, fill=NAVY, align=CEN)
    _c(ws, 2, 3, LB["detail_header"], font=hdr, fill=NAVY, align=CEN)
    ws.merge_cells(start_row=2, start_column=EB0, end_row=2, end_column=EB1)
    _c(ws, 2, EB0, LB["evidence_before"], font=hdr, fill=NAVY, align=CEN)
    ws.merge_cells(start_row=2, start_column=EA0, end_row=2, end_column=EA1)
    _c(ws, 2, EA0, LB["evidence_after"], font=hdr, fill=NAVY, align=CEN)

    r = 3
    for tc in tcs:
        b = r
        # -- header row (TC-id + title + priority) --
        _c(ws, b, 1, tc.get("id"), font=Font(bold=True, color=WHITE), fill=BLUE2, align=CEN)
        ws.merge_cells(start_row=b, start_column=2, end_row=b, end_column=3)
        _c(ws, b, 2, f"[{tc.get('screen', '')}]  {tc.get('title', '')}",
           font=Font(bold=True, color=WHITE), fill=BLUE2)
        pri = tc.get("pri", "")
        ws.merge_cells(start_row=b, start_column=EB0, end_row=b, end_column=LAST_COL)
        _c(ws, b, EB0, f"Priority: {pri}", font=Font(bold=True, color=WHITE),
           fill=PRI.get(pri, PRI["_default"]), align=Alignment(vertical="center"))
        # -- field rows --
        for i, (num, label, key) in enumerate(FIELDS, 1):
            rr = b + i
            _c(ws, rr, 1, num, fill=NUMF, align=CEN)
            _c(ws, rr, 2, label, font=Font(bold=True), fill=LABEL)
            cell = _c(ws, rr, 3, tc.get(key, ""))
            if key == "actual":
                res = tc.get("result")
                if res in RES_FILL:
                    cell.fill = _f(RES_FILL[res])
                    cell.font = Font(bold=True, color=RES_FONT[res])
        # -- note row --
        nrow = b + 6
        _c(ws, nrow, 1, LB["note_icon"], fill=NOTEF, align=CEN)
        _c(ws, nrow, 2, LB["note"], font=Font(bold=True), fill=NOTEF)
        _c(ws, nrow, 3, tc.get("note") or "", fill=NOTEF)
        # -- evidence merged region: fields+note (b+1 .. b+6) --
        ev_top, ev_bot = b + 1, b + 6
        heights = []
        for col_start, col_end, key in ((EB0, EB1, "before"), (EA0, EA1, "after")):
            ws.merge_cells(start_row=ev_top, start_column=col_start, end_row=ev_bot, end_column=col_end)
            fname = tc.get(key)
            p = (shots_dir / fname) if fname else None
            anchor = _c(ws, ev_top, col_start, None, align=CEN)
            if p and p.exists() and PILImage is not None:
                w, h = _img_size(p)
                img = XLImage(str(p))
                img.width, img.height = w, h
                ws.add_image(img, f"{get_column_letter(col_start)}{ev_top}")
                heights.append(h)
            elif fname:
                anchor.value = LB["placeholder_missing"].format(name=fname)
            else:
                anchor.value = LB["placeholder_na"]
                anchor.fill = _f(NAF)
        # -- row heights: chia đủ cao cho ảnh (max before/after) --
        img_h = max(heights) if heights else 0
        if img_h:
            per = max(LY["evidence_row_min_height"], (img_h * 0.75) / 6.0)  # px→pt, chia 6 dòng
            for rr in range(ev_top, ev_bot + 1):
                ws.row_dimensions[rr].height = per
        else:
            for rr in range(ev_top, ev_bot + 1):
                ws.row_dimensions[rr].height = LY["evidence_row_empty_height"]
        r = b + 8  # 7 dòng block + 1 spacer

    ws.freeze_panes = "A3"


# ---------------- Checklist ----------------
def build_checklist(wb, tcs):
    ws = wb.create_sheet("Checklist")
    for col, w in LY["checklist_cols"].items():
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:F1")
    _c(ws, 1, 1, LB["checklist_title"], font=Font(bold=True, size=FT["checklist_title"], color=WHITE),
       fill=NAVY, align=Alignment(vertical="center"))
    hdr = Font(bold=True, color=WHITE)
    for c, name in enumerate(LB["checklist_header"], 1):
        _c(ws, 3, c, name, font=hdr, fill=BLUE2, align=CEN)
    r = 4
    for i, tc in enumerate(tcs, 1):
        row_fill = WHITE if i % 2 else ALT
        _c(ws, r, 1, i, fill=row_fill, align=CEN)
        _c(ws, r, 2, tc.get("id"), fill=row_fill)
        _c(ws, r, 3, f"[{tc.get('screen', '')}] {tc.get('title', '')}", fill=row_fill)
        _c(ws, r, 4, tc.get("pri"), fill=row_fill, align=CEN)
        res = tc.get("result", "")
        rc = _c(ws, r, 5, res, fill=RES_FILL.get(res, row_fill), align=CEN)
        if res in RES_FONT:
            rc.font = Font(bold=True, color=RES_FONT[res])
        _c(ws, r, 6, tc.get("source", ""), fill=row_fill)
        r += 1
    last = r - 1
    # hàng tổng: mỗi status 1 cặp (nhãn, COUNTIF) — lấy thẳng từ theme
    _c(ws, r, 1, "Total", font=Font(bold=True), fill=LABEL)
    _c(ws, r, 2, f"=COUNTA(B4:B{last})", font=Font(bold=True), fill=LABEL, align=CEN)
    col = 3
    for name, spec in LB["status"].items():
        fill = P[spec["bg"]]
        _c(ws, r, col, name, fill=fill, align=CEN)
        _c(ws, r, col + 1, f'=COUNTIF(E4:E{last},"{spec["count_match"]}")', fill=fill, align=CEN)
        col += 2
        if col > 6:  # tràn sang hàng dưới nếu quá 6 cột
            r += 1
            col = 3


def main():
    if len(sys.argv) < 3:
        print("usage: build_evidence.py <tcs.json> <out.xlsx>")
        sys.exit(1)
    tcs_path, out = Path(sys.argv[1]), Path(sys.argv[2])
    data = json.loads(tcs_path.read_text(encoding="utf-8"))
    meta, tcs = data.get("meta", {}), data.get("tcs", [])
    shots_dir = tcs_path.parent / data.get("shots_dir", "shots")

    unknown = {t.get("result") for t in tcs} - set(LB["status"]) - {None}
    if unknown:
        print(f"WARN: result lạ (không có trong theme.json): {unknown}", file=sys.stderr)

    wb = Workbook()
    build_cover(wb, meta)
    build_testcases(wb, meta, tcs, shots_dir)
    build_checklist(wb, tcs)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    gaps = sum(1 for t in tcs if t.get("result") == "SPEC-GAP")
    extra = f", {gaps} SPEC-GAP" if gaps else ""
    print(f"OK -> {out}  ({len(tcs)} cases{extra}, 3 sheets: Cover/Test Cases/Checklist)")


if __name__ == "__main__":
    main()
